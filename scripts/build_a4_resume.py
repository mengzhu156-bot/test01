from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageOps, ImageDraw


ROOT = Path("/Users/mengzhu/Documents/个人网站")
OUT = ROOT / "output" / "resume" / "孟竹_AI美术导演_A4竖版简历.docx"
PORTRAIT = ROOT / "src" / "assets" / "profile-portrait.jpg"
TMP_PHOTO = ROOT / "tmp" / "resume_portrait_bw.png"

FONT_CN = "Microsoft YaHei"
FONT_EN = "Arial"
INK = "111111"
MUTED = "666666"
LIGHT = "E7E7E7"
ACCENT = "F2B600"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            edge_data = kwargs.get(edge)
            tag = "w:{}".format(edge)
            element = borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                borders.append(element)
            for key in ["sz", "val", "color", "space"]:
                if key in edge_data:
                    element.set(qn("w:{}".format(key)), str(edge_data[key]))


def set_cell_margins(cell, top=40, start=70, bottom=40, end=70):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width
            tc_pr = row.cells[idx]._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width.inches * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def strip_table_borders(table):
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(
                cell,
                top={"val": "nil"},
                bottom={"val": "nil"},
                left={"val": "nil"},
                right={"val": "nil"},
            )


def style_run(run, size=9, bold=False, color=INK, font=FONT_CN, italic=False):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(paragraph, text, size=9, bold=False, color=INK, font=FONT_CN, italic=False):
    run = paragraph.add_run(text)
    style_run(run, size=size, bold=bold, color=color, font=font, italic=italic)
    return run


def set_para(paragraph, before=0, after=0, line=1.05, align=None):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    if align is not None:
        paragraph.alignment = align


def add_heading(cell, title):
    p = cell.add_paragraph()
    set_para(p, before=0, after=7)
    add_text(p, title.upper(), 11.5, True, INK, FONT_EN)


def add_kicker(cell, title):
    p = cell.add_paragraph()
    set_para(p, before=0, after=4)
    add_text(p, f"// {title}", 8.2, True, INK, FONT_EN)


def bullet(cell, text, size=7.2, after=2.2):
    p = cell.add_paragraph()
    set_para(p, after=after, line=1.08)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.first_line_indent = Inches(-0.08)
    add_text(p, "• ", size, False, INK, FONT_EN)
    add_text(p, text, size, False, MUTED)
    return p


def body_line(cell, text, size=7.5, bold=False, color=MUTED, after=1.8):
    p = cell.add_paragraph()
    set_para(p, after=after, line=1.05)
    add_text(p, text, size, bold, color)
    return p


def make_photo():
    if not PORTRAIT.exists():
        return None
    img = Image.open(PORTRAIT).convert("RGB")
    w, h = img.size
    target_ratio = 0.78
    crop_h = h
    crop_w = int(crop_h * target_ratio)
    if crop_w > w:
        crop_w = w
        crop_h = int(crop_w / target_ratio)
    left = max(0, int((w - crop_w) * 0.5))
    top = max(0, int(h * 0.05))
    img = img.crop((left, top, left + crop_w, top + crop_h))
    img = ImageOps.grayscale(img)
    img = ImageOps.autocontrast(img)
    img = ImageOps.colorize(img, black="#080808", white="#f1f1f1")
    img = img.resize((720, 920))
    canvas = Image.new("RGB", (720, 920), "#eeeeee")
    canvas.paste(img, (0, 0))
    draw = ImageDraw.Draw(canvas)
    draw.polygon([(0, 850), (74, 920), (0, 920)], fill=f"#{ACCENT}")
    TMP_PHOTO.parent.mkdir(exist_ok=True)
    canvas.save(TMP_PHOTO)
    return TMP_PHOTO


def add_metric_bar(cell, label, dots=4):
    p = cell.add_paragraph()
    set_para(p, after=2, line=1)
    add_text(p, label, 7.0, False, MUTED)
    add_text(p, "  ", 7)
    for idx in range(5):
        add_text(p, "●", 6.5, False, INK if idx < dots else "CFCFCF", FONT_EN)


def build():
    photo = make_photo()
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.0)
    section.bottom_margin = Cm(0.9)
    section.left_margin = Cm(1.05)
    section.right_margin = Cm(1.05)
    section.header_distance = Cm(0.4)
    section.footer_distance = Cm(0.4)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_CN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal.font.size = Pt(8)
    normal.paragraph_format.space_after = Pt(0)

    page = doc.add_table(rows=1, cols=2)
    page.autofit = False
    strip_table_borders(page)
    set_table_width(page, [Inches(2.25), Inches(4.93)])
    left, right = page.rows[0].cells
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_margins(left, top=0, start=0, bottom=0, end=120)
    set_cell_margins(right, top=0, start=120, bottom=0, end=0)

    if photo:
        p = left.paragraphs[0]
        set_para(p, after=12)
        r = p.add_run()
        r.add_picture(str(photo), width=Inches(1.92), height=Inches(2.45))

    add_heading(left, "Skills")
    add_kicker(left, "AIGC / TECH")
    for item in [
        "Midjourney / Stable Diffusion",
        "Runway / Sora / Veo3",
        "关键帧 / 风格迭代",
        "角色一致性 / 场景连贯性",
    ]:
        bullet(left, item)
    add_kicker(left, "DESIGN")
    for item in [
        "角色与场景概念设定",
        "色彩体系 / 光影氛围",
        "Photoshop / Illustrator / InDesign",
        "Blender / SketchUp / AutoCAD",
    ]:
        bullet(left, item)

    add_heading(left, "Education")
    body_line(left, "上海戏剧学院", 7.7, True, INK)
    body_line(left, "艺术与科技 硕士 | 2023 - 2026", 7.2)
    body_line(left, "南京大学金陵学院", 7.7, True, INK, after=0.8)
    body_line(left, "环境设计 本科 | GPA 4.17", 7.2)
    body_line(left, "台湾元智大学交换 / 日本语进修", 7.0)

    add_heading(left, "Honors")
    for item in [
        "2026 香港国际 AI 电影节优秀作品入选《重生》",
        "2024 大学生 AI 艺术季 AI 影像创作单元入围",
        "平面设计师证 / 美术教师资格证",
        "英语 CET-6 工作语言 / 日语 N2",
    ]:
        bullet(left, item, 7.0)

    add_heading(left, "Contact")
    for item in ["P: 15605162339", "E: 740050785@qq.com", "A: 上海", "D: AI 美术 / AIGC 视觉指导"]:
        body_line(left, item, 7.3, False, MUTED)

    header = right.add_table(rows=1, cols=2)
    header.autofit = False
    strip_table_borders(header)
    set_table_width(header, [Inches(3.05), Inches(1.7)])
    name_cell, contact_cell = header.rows[0].cells
    set_cell_margins(name_cell, top=0, start=0, bottom=0, end=50)
    set_cell_margins(contact_cell, top=16, start=40, bottom=0, end=0)

    p = name_cell.paragraphs[0]
    set_para(p, after=1)
    add_text(p, "/", 27, True, ACCENT, FONT_EN)
    add_text(p, " Meng", 20, True, INK, FONT_EN)
    p = name_cell.add_paragraph()
    set_para(p, after=2)
    add_text(p, "ZHU", 30, True, INK, FONT_EN)
    p = name_cell.add_paragraph()
    set_para(p, after=8)
    add_text(p, "│  AI ART DIRECTOR / AIGC VISUAL LEAD", 9.5, False, INK, FONT_EN)

    for label, value in [
        ("P", "15605162339"),
        ("E", "740050785@qq.com"),
        ("L", "Shanghai"),
        ("S", "Visual / Brand / Spatial Design"),
    ]:
        p = contact_cell.add_paragraph()
        set_para(p, after=1.4)
        add_text(p, f"{label}: ", 6.8, True, INK, FONT_EN)
        add_text(p, value, 6.8, False, MUTED, FONT_EN)

    p = right.add_paragraph()
    set_para(p, before=4, after=11, line=1.18)
    add_text(
        p,
        "艺术设计硕士，AIGC 视觉与美术方向深度实践者。擅长以导演思维统领 AI 影像叙事，从概念设定、关键帧、角色与场景构建到成片审校，稳定控制风格一致性、情绪氛围与商业落地品质。",
        8.0,
        False,
        MUTED,
    )

    add_heading(right, "Work Experience")
    exp = [
        (
            "漫剧项目",
            "AI 美术主导 / 视频质量把控 | 2025.08 - 2026.03",
            [
                "负责整体视觉风格、美术方向、角色设计、场景构建与画面氛围统一。",
                "主导 AI 画面生成与成片视觉审核，保障角色一致性、场景连贯性与剪辑节奏。",
                "作品于红果上线，斩获千万级热度，验证 AI 内容商业化转化能力。",
            ],
        ),
        (
            "APB 艺术节 & 无锡蠡湖生态设计节",
            "视觉设计师 | 2024.03 - 2024.12",
            [
                "输出 AIGC 展演视觉方案，落地主视觉海报、KV 展板、场刊与展陈物料。",
                "从 0 到 1 参与生态设计节摄影展策划、建模与现场布置，观展 10 万+，媒介报道 10+。",
            ],
        ),
        (
            "上海铁海文化艺术交流策划有限公司",
            "设计师（实习） | 2023.09 - 2023.11",
            [
                "承接西岸艺术与博览会展位空间设计，以用户体验思维完成功能分区和模块化布局。",
                "完成展签、画册、施工深化与现场布展协作，推动展览按期落地。",
            ],
        ),
    ]

    for title, subtitle, points in exp:
        row = right.add_table(rows=1, cols=2)
        row.autofit = False
        strip_table_borders(row)
        set_table_width(row, [Inches(0.18), Inches(4.55)])
        line_cell, text_cell = row.rows[0].cells
        set_cell_margins(line_cell, top=0, start=0, bottom=0, end=0)
        set_cell_margins(text_cell, top=0, start=50, bottom=5, end=0)
        set_cell_border(line_cell, right={"val": "single", "sz": "8", "color": "555555"})
        p = text_cell.paragraphs[0]
        set_para(p, after=1.2)
        add_text(p, title, 9.5, True, INK)
        p = text_cell.add_paragraph()
        set_para(p, after=2.2)
        add_text(p, subtitle, 7.4, False, INK, italic=True)
        for point in points:
            bullet(text_cell, point, 7.1, after=1.1)

    add_heading(right, "Core AIGC Projects")
    projects = [
        ("AI 影片《重生》", "导演 / AI 美术 / 视觉 | 2026", "入围首届香港国际 AI 电影节官方评选；完成非线性叙事结构、整体视觉风格、角色与场景构建，并把控全片风格一致性。"),
        ("生成式 AI 艺术展《创世纪 · 和合共生》", "AI 视觉主创（影像 + 装置） | 2024 - 2025", "负责整体视觉语言、关键帧生成与互动装置设计；项目于上海展出并赴保加利亚巡展。"),
        ("沉浸式 AI&MR 戏曲《黛玉葬花》", "AIGC 视觉创意 | 2024", "完成园林微观场景与 MR 的视觉融合，支撑戏曲叙事的情绪与氛围呈现。"),
        ("未来戏剧展厅", "AI 生成视觉内容 / UI 设计", "负责 AI 生成音乐区域 UI 与视觉内容，落地图形智能编译、自然语言驱动动态视觉生成亮点。"),
    ]
    for title, subtitle, desc in projects:
        p = right.add_paragraph()
        set_para(p, after=0.8)
        add_text(p, title, 8.2, True, INK)
        add_text(p, f"  {subtitle}", 7.0, False, MUTED)
        p = right.add_paragraph()
        set_para(p, after=3.2, line=1.08)
        add_text(p, desc, 7.0, False, MUTED)

    footer = doc.add_table(rows=1, cols=4)
    footer.autofit = False
    strip_table_borders(footer)
    set_table_width(footer, [Inches(1.75), Inches(1.75), Inches(1.75), Inches(1.75)])
    labels = [
        ("AIGC VISUAL", [("Style Control", 5), ("Keyframes", 5), ("Video Tools", 4)]),
        ("ART DIRECTION", [("Character", 5), ("Scene", 4), ("Color / Light", 5)]),
        ("SPATIAL DESIGN", [("Exhibition", 5), ("CAD / SketchUp", 4), ("Execution", 4)]),
        ("BRAND TOOLS", [("PS / AI / ID", 5), ("Blender", 4), ("Bilingual", 4)]),
    ]
    for idx, (heading, bars) in enumerate(labels):
        cell = footer.rows[0].cells[idx]
        set_cell_margins(cell, top=80, start=70, bottom=0, end=70)
        set_cell_border(cell, top={"val": "single", "sz": "6", "color": LIGHT})
        p = cell.paragraphs[0]
        set_para(p, after=5)
        add_text(p, heading, 9.0, True, INK, FONT_EN)
        for label, dots in bars:
            add_metric_bar(cell, label, dots)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
